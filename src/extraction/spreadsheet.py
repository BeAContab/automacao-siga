from __future__ import annotations

"""Leitura da planilha XLSX com os CNPJs usados pela automação."""

import logging
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.workbook import Workbook

from src.utils.text import strip_accents


LOGGER = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Constantes de cabeçalho — usadas tanto na leitura quanto na escrita de status.
# Manter em um único lugar evita divergência entre as duas operações.
# ---------------------------------------------------------------------------
_REQUIRED_DOCUMENT_HEADERS: frozenset[str] = frozenset({"cnpj", "cgf"})
_COD_HEADERS: frozenset[str] = frozenset({"cod", "codigo", "código"})
_COMPANY_HEADERS: frozenset[str] = frozenset({
    "empresa", "nome empresa", "razao social", "razão social", "razao social completa"
})
_ALL_KNOWN_HEADERS: frozenset[str] = _REQUIRED_DOCUMENT_HEADERS | _COD_HEADERS | _COMPANY_HEADERS

# Mantém o workbook de status aberto em memória por caminho de planilha durante um lote,
# evitando reabrir e reescanear o cabeçalho a cada célula (ver PLANO_CORRECOES.md item 7).
_STATUS_WORKBOOK_CACHE: dict[Path, tuple[Workbook, dict[str, int]]] = {}


@dataclass(slots=True)
class SpreadsheetRow:
    """Representa uma linha válida da planilha já normalizada para processamento."""
    row_number: int
    cnpj: str
    cod: str = ""
    empresa: str = ""


def load_cnpjs_from_xlsx(path: Path) -> list[SpreadsheetRow]:
    """Localiza as colunas da planilha, normaliza os valores e devolve apenas linhas válidas."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active

    # Varredura de cabeçalho via função compartilhada (ver item 22 do PLANO_CORRECOES.md)
    header_row_number, header_map = _scan_header_row(sheet)

    if "cnpj" not in header_map and "cgf" not in header_map:
        workbook.close()
        raise ValueError("A planilha precisa conter uma coluna chamada 'cnpj'.")

    document_index = header_map.get("cnpj", header_map.get("cgf"))
    cod_index = _find_header_index(header_map, _COD_HEADERS)
    company_index = _find_header_index(header_map, _COMPANY_HEADERS)
    rows: list[SpreadsheetRow] = []
    # Cada CNPJ é lido já sem caracteres de formatação para bater com a forma canônica usada no SIGA.
    for row_number, row in enumerate(
        sheet.iter_rows(min_row=header_row_number + 1, values_only=True),
        start=header_row_number + 1,
    ):
        value = row[document_index] if document_index is not None and document_index < len(row) else None
        digits = _normalize_document(value)
        if not digits:
            continue
        # Rejeita documentos com formato obviamente inválido antes de tentar a busca no SIGA.
        if not _is_likely_valid_document(digits):
            LOGGER.warning(
                "Linha %s ignorada: o valor '%s' nao parece um CNPJ/CGF valido (verifique a planilha).",
                row_number,
                value,
            )
            continue
        cod_value = _normalize_free_text(row[cod_index]) if cod_index is not None and cod_index < len(row) else ""
        company_value = _normalize_free_text(row[company_index]) if company_index is not None and company_index < len(row) else ""
        rows.append(
            SpreadsheetRow(
                row_number=row_number,
                cnpj=digits,
                cod=cod_value,
                empresa=company_value,
            )
        )

    workbook.close()

    if not rows:
        raise ValueError("Nenhum CNPJ valido foi encontrado na planilha informada.")

    return rows


def load_cod_empresa_cnpj_base(path: Path) -> dict[str, SpreadsheetRow]:
    """Lê a planilha-base COD/EMPRESA/CNPJ (mesmas colunas de `load_cnpjs_from_xlsx`,
    por isso reaproveitada aqui) e devolve um mapa cnpj -> linha.

    Usada como rede de segurança para nomear a pasta de saída (`<COD - EMPRESA - CNPJ>`)
    quando não há como inferir isso a partir da estrutura de pastas de entrada — ex.:
    modo NF-e recebendo um arquivo `.xlsx` solto, fora da convenção de pastas do SIGA.
    Nunca lança exceção: pasta ausente ou planilha ilegível apenas resultam num mapa
    vazio, e quem chama cai no fallback "SEM-COD"/"SEM-EMPRESA" já existente.
    """
    if not path.exists():
        return {}
    try:
        rows = load_cnpjs_from_xlsx(path)
    except Exception:  # noqa: BLE001
        LOGGER.exception("Falha ao ler a planilha-base COD/EMPRESA/CNPJ %s", path)
        return {}
    return {row.cnpj: row for row in rows}


# Pastas de saída (Emissor) e entrada (Destinatário) — mesmos rótulos para os modos NF-e
# e NFC-e, para o escritório reconhecer a convenção em qualquer um dos dois.
DIRECAO_SUBPASTA: dict[str, str] = {"saida": "Notas de Saída", "entrada": "Notas de Entrada"}


def tipo_nota_do_nome_arquivo(nome_arquivo: str) -> str:
    """Detecta se um relatório de detalhamento é de notas de SAÍDA (perfil "Emissor" no
    SIGA — a empresa emitiu o documento) ou de ENTRADA (perfil "Destinatario" — a
    empresa recebeu), pelo nome do arquivo.

    A convenção vem do próprio SIGA: `_display_label` (`siga_extractor.py`) grava o
    rótulo do perfil ("Emissor"/"Destinatario") como parte literal do nome do arquivo de
    detalhamento gerado (ex.: "Informações Fiscais - NFC-e - Emissor - Detalhamento
    Agosto de 2026...xlsx"). Devolve "" quando o nome não segue essa convenção (arquivo
    renomeado manualmente, ou fora do padrão do SIGA) — quem chama decide o fallback
    nesse caso (hoje: não separar em subpasta, mantendo o comportamento anterior).
    """
    normalizado = strip_accents(nome_arquivo).lower()
    if "emissor" in normalizado:
        return "saida"
    if "destinatario" in normalizado:
        return "entrada"
    return ""


def load_cnpjs_from_text(text: str) -> list[SpreadsheetRow]:
    """Lê CNPJs informados manualmente, um por linha ou separados por pontuação."""
    rows: list[SpreadsheetRow] = []
    seen: set[str] = set()

    # A numeração começa em 2 (não em 1) para reservar a linha 1 para o cabeçalho da
    # planilha de resultados gerada a partir da entrada manual (ver build_manual_entry_workbook).
    for line_number, raw_line in enumerate(text.splitlines(), start=2):
        # A interface permite colar CNPJs em vários formatos; tudo é normalizado para 14 dígitos.
        normalized_values = _extract_documents_from_text(raw_line)
        for value in normalized_values:
            if value in seen:
                continue
            # Rejeita documentos com formato obviamente inválido antes de tentar a busca no SIGA.
            if not _is_likely_valid_document(value):
                LOGGER.warning(
                    "Linha %s ignorada: o valor '%s' nao parece um CNPJ/CGF valido.",
                    line_number,
                    value,
                )
                continue
            seen.add(value)
            rows.append(SpreadsheetRow(row_number=line_number, cnpj=value))

    if not rows:
        raise ValueError("Nenhum CNPJ valido foi informado manualmente.")

    return rows


def build_manual_entry_workbook(rows: list[SpreadsheetRow], destination_path: Path) -> None:
    """Gera uma planilha de resultados para CNPJs digitados manualmente na GUI.

    Sem esta planilha, a entrada manual não deixaria nenhum rastro persistido de status
    além do log em tela, que se perde ao fechar a aplicação.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SIGA EMPS"
    headers = ["cnpj", "NF-e", "NFC-e", "CT-e", "Malha Fiscal", "Débitos Fiscais"]
    for column_index, header in enumerate(headers, start=1):
        sheet.cell(row=1, column=column_index, value=header)
    for row in rows:
        sheet.cell(row=row.row_number, column=1, value=row.cnpj)
    destination_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination_path)
    workbook.close()


def _normalize_header(value: object) -> str:
    """Remove acentos e padroniza o cabeçalho para comparação segura."""
    if value is None:
        return ""
    return strip_accents(str(value)).strip().lower()


def _normalize_document(value: object) -> str:
    # Normaliza o documento preservando letras e números para dar suporte ao CNPJ alfanumérico.
    if value is None:
        return ""
    text = "".join(char for char in str(value) if char.isalnum())
    # O CNPJ precisa permanecer com 14 caracteres, inclusive os zeros à esquerda se for numérico.
    if not text:
        return ""
    return text.zfill(14) if len(text) <= 14 else text


def _is_likely_valid_document(digits: str) -> bool:
    """Verifica se um documento normalizado tem formato minimamente aceitável.

    Não implementa validação de dígito verificador (incompatível com CNPJ alfanumérico).
    Rejeita apenas casos claramente inválidos:
    - Comprimento diferente de 14 caracteres.
    - Sequências puramente numéricas com todos os dígitos iguais (ex: 00000000000000).
    """
    if len(digits) != 14:
        return False
    # Rejeição de sequências trivialmente inválidas (apenas para documentos numéricos)
    if digits.isdigit() and len(set(digits)) == 1:
        return False
    return True


def _normalize_free_text(value: object) -> str:
    """Preserva o texto útil da célula e remove apenas espaços excedentes."""
    if value is None:
        return ""
    return str(value).strip()


def _find_header_index(header_map: dict[str, int], aliases: frozenset[str]) -> int | None:
    """Localiza o primeiro índice disponível para um grupo de aliases equivalentes."""
    for alias in aliases:
        index = header_map.get(alias)
        if index is not None:
            return index
    return None


def _scan_header_row(sheet) -> tuple[int, dict[str, int]]:
    """Varre as linhas da planilha em busca do cabeçalho e devolve (row_number, header_map).

    Função compartilhada entre leitura e escrita de status, eliminando a duplicação
    anterior (ver PLANO_CORRECOES.md item 22). O índice retornado é 0-based (como usado
    em iter_rows); a camada de escrita de status converte para 1-based ao precisar.
    """
    header_map: dict[str, int] = {}
    header_row_number = 0
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        normalized_cells = [_normalize_header(cell) for cell in row]
        if any(cell in normalized_cells for cell in _ALL_KNOWN_HEADERS):
            header_row_number = row_number
            for index, value in enumerate(normalized_cells):
                if value:
                    header_map[value] = index
            break
    return header_row_number, header_map


def _extract_documents_from_text(value: object) -> list[str]:
    """Extrai todos os blocos alfanuméricos de 14 caracteres presentes em um texto livre."""
    if value is None:
        return []

    text = str(value)
    matches = []
    for token in text.replace("\t", " ").replace(";", " ").replace(",", " ").split():
        digits = _normalize_document(token)
        if len(digits) == 14:
            matches.append(digits)

    if matches:
        return matches

    digits = _normalize_document(text)
    return [digits] if len(digits) == 14 else []


def _build_status_header_map(sheet) -> dict[str, int]:
    """Localiza a linha de cabeçalho da planilha de status (usa _scan_header_row compartilhada).

    Retorna header_map com índices 1-based (para openpyxl), convertendo a partir dos
    índices 0-based retornados por _scan_header_row.
    """
    _, header_map_zero = _scan_header_row(sheet)
    # Converter para 1-based (padrão do openpyxl para escrita de células)
    return {key: val + 1 for key, val in header_map_zero.items()}


def _open_status_workbook(spreadsheet_path: Path) -> tuple[Workbook, dict[str, int]] | None:
    """Abre (ou reaproveita do cache) o workbook de status para um caminho de planilha."""
    cached = _STATUS_WORKBOOK_CACHE.get(spreadsheet_path)
    if cached is not None:
        return cached

    try:
        # Usar data_only=False para preservar as fórmulas originais da planilha
        workbook = load_workbook(spreadsheet_path, read_only=False, data_only=False)
    except (PermissionError, OSError):
        LOGGER.warning(
            "Não foi possível abrir a planilha de resultados '%s' para gravar o status (arquivo pode estar aberto no Excel).",
            spreadsheet_path,
        )
        return None

    sheet = workbook["SIGA EMPS"] if "SIGA EMPS" in workbook.sheetnames else workbook.active
    header_map = _build_status_header_map(sheet)
    cached = (workbook, header_map)
    _STATUS_WORKBOOK_CACHE[spreadsheet_path] = cached
    return cached


def close_status_workbook(spreadsheet_path: Path | None) -> None:
    """Fecha e descarta do cache o workbook de status aberto para o caminho informado.

    Deve ser chamado ao final de cada lote para liberar o arquivo e evitar que um
    workbook de uma execução anterior seja reaproveitado com cabeçalhos desatualizados.
    """
    if not spreadsheet_path:
        return
    cached = _STATUS_WORKBOOK_CACHE.pop(spreadsheet_path, None)
    if cached is not None:
        workbook, _header_map = cached
        try:
            workbook.close()
        except Exception:  # noqa: BLE001
            LOGGER.warning("Falha ao fechar a planilha de resultados '%s'.", spreadsheet_path)


def write_status_to_spreadsheet_cell(
    spreadsheet_path: Path,
    row_number: int,
    document_tab: str,
    status_str: str,
) -> None:
    """Escreve em tempo real o status do download na coluna correspondente na planilha de resultados.

    O workbook fica aberto em memória (ver `_STATUS_WORKBOOK_CACHE`) durante todo o lote, para que
    cada chamada só precise gravar a célula e salvar, sem reabrir e reescanear o arquivo inteiro.
    """
    if not spreadsheet_path or not row_number:
        return

    cached = _open_status_workbook(spreadsheet_path)
    if cached is None:
        return
    workbook, header_map = cached

    sheet = workbook["SIGA EMPS"] if "SIGA EMPS" in workbook.sheetnames else workbook.active

    # Normalizar o nome da aba fiscal para fazer a correspondência com a coluna
    norm_tab = _normalize_header(document_tab)
    col_index = None

    # Mapear a coluna que contém o nome da aba fiscal (ex: "nf-e" ou "debitos fiscais")
    for key, val in header_map.items():
        if norm_tab in key or key in norm_tab:
            col_index = val
            break

    if col_index is None:
        LOGGER.warning(
            "Não foi possível localizar a coluna da aba '%s' na planilha de resultados '%s'; status '%s' não gravado.",
            document_tab,
            spreadsheet_path,
            status_str,
        )
        return

    sheet.cell(row=row_number, column=col_index, value=status_str)
    try:
        workbook.save(spreadsheet_path)
    except (PermissionError, OSError):
        LOGGER.warning(
            "Não foi possível salvar a planilha de resultados '%s' agora (arquivo pode estar aberto no Excel). "
            "O status '%s' permanece pendente em memória e será regravado na próxima atualização.",
            spreadsheet_path,
            status_str,
        )

