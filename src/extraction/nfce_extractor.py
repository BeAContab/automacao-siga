from __future__ import annotations

"""Extração de XML de NFC-e via login automático no portal da SEFAZ-CE.

Porta a lógica de negócio de `importação/NFCE/codigo-fonte/xml nfce.py` (login,
gestão de sessão, download por chave de acesso de 44 dígitos) para usar
`BrowserSession`/`selenium_compat.py` em vez do gerenciamento próprio de
Chrome/driver do script original — o próprio `GerenciadorSessao.iniciar_chrome`/
`conectar_selenium` daquele arquivo é redundante com `src/utils/browser.py`.

O arquivo-fonte original nunca é importado ou copiado para dentro deste módulo:
ele contém um CPF e uma senha de contador em texto puro (hardcoded), então toda
portagem é reescrita manual — aqui as credenciais só vêm de `Settings.nfce_cpf`/
`Settings.nfce_senha`, preenchidos em runtime a partir de campos na GUI
(`src/gui.py`, `nfce_cpf_var`/`nfce_senha_var`) e nunca persistidos em disco.
"""

from dataclasses import dataclass
from pathlib import Path
import logging
import re
import shutil
import threading
import time

from openpyxl import load_workbook

from src.config import Settings
from src.extraction.spreadsheet import (
    DIRECAO_SUBPASTA,
    SpreadsheetRow,
    load_cod_empresa_cnpj_base,
    tipo_nota_do_nome_arquivo,
)
from src.utils.execution_control import wait_if_paused
from src.utils.narration import narrate, narrate_error, narrate_success, narrate_warning
from src.utils.selenium_compat import BrowserContext, Error, Page, TimeoutError, UnexpectedAlertError

LOGGER = logging.getLogger(__name__)


def _only_digits(value: str) -> str:
    return re.sub(r"\D", "", str(value or ""))


def normalize_cnpj(value: str) -> str:
    digits = _only_digits(value)
    return digits.zfill(14)[-14:] if digits else ""


def normalize_ie(value: str) -> str:
    text = str(value or "").strip().replace("\xa0", " ")
    upper = text.upper()
    if not text or "S/ INSCRI" in upper or "SEM INSCRI" in upper:
        return ""
    return _only_digits(text).lstrip("0")


def sanitize_folder_name(value: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*]', "", value)
    cleaned = re.sub(r"\s+", " ", cleaned).strip().rstrip(".")
    return cleaned or "SEM-NOME"


def _is_nfce_key(digits: str) -> bool:
    """True se a sequência é uma chave de acesso de 44 dígitos do modelo NFC-e (65).

    Posições 21-22 (índice [20:22]) da chave de acesso identificam o modelo do
    documento fiscal no padrão nacional: 65 = NFC-e, 55 = NF-e, 57 = CT-e. NF-e, NFC-e
    e CT-e usam o mesmo formato de 44 dígitos, então essa checagem é a rede de
    segurança que impede uma chave de NF-e/CT-e (ex.: achada por engano numa pasta
    'NF-e'/'CT-e' vizinha) de ser tratada como chave de NFC-e.
    """
    return len(digits) == 44 and digits[20:22] == "65"


def is_nfe_key(digits: str) -> bool:
    """True se a sequência é uma chave de acesso de 44 dígitos do modelo NF-e (55).

    Mesma rede de segurança que `_is_nfce_key`, só que para o modelo NF-e — usada pelo
    modo NF-e (`meudanfe_extractor.py`) para descartar uma chave de NFC-e/CT-e que
    porventura apareça na coluna "Chave NF-e" (dado errado na planilha de origem).
    """
    return len(digits) == 44 and digits[20:22] == "55"


def cnpj_from_access_key(digits: str) -> str:
    """CNPJ do emitente embutido na própria chave de acesso (posições 7 a 20).

    Padrão nacional de 44 dígitos: UF(2) + AAMM(4) + CNPJ(14) + modelo(2) + série(3) +
    número(9) + tpEmis(1) + cNF(8) + DV(1) — o mesmo para NF-e, NFC-e e CT-e. Permite
    identificar de qual empresa é uma chave sem depender de nome de arquivo/pasta —
    útil quando as chaves vêm de um arquivo solto, fora da convenção de pastas do SIGA.
    """
    return digits[6:20] if len(digits) == 44 else ""


def extract_keys_from_spreadsheet(path: Path) -> list[str]:
    """Varre todas as colunas/linhas em busca de chaves de acesso de NFC-e (44 dígitos, modelo 65)."""
    keys: list[str] = []
    seen: set[str] = set()
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        for row in sheet.iter_rows(values_only=True):
            for cell in row:
                if cell is None:
                    continue
                digits = _only_digits(str(cell))
                if _is_nfce_key(digits) and digits not in seen:
                    seen.add(digits)
                    keys.append(digits)
        workbook.close()
    except Exception:  # noqa: BLE001
        LOGGER.exception("Falha ao ler chaves da planilha %s", path)
    return keys


def _company_folder_for_cnpj(keys_folder: Path, cnpj: str) -> Path | None:
    """Pasta da empresa dentro de `keys_folder`, cujo nome contém o CNPJ (convenção do
    modo SIGA: 'COD - EMPRESA - CNPJ', ver `_build_taxpayer_folder_name` em
    `siga_extractor.py`) — aceita `keys_folder` apontando tanto para a pasta raiz de
    saída do SIGA (o CNPJ está numa subpasta direta) quanto para a pasta de uma empresa
    específica selecionada diretamente (o CNPJ está no próprio nome de `keys_folder`).

    Não vai mais fundo que isso de propósito: numa pasta de nível inferior (ex. o mês
    ou a subpasta "NFC-e" em si) não sobra nenhum sinal do CNPJ no caminho, então
    "adivinhar" a que empresa ela pertence arriscaria atribuir documentos de uma
    empresa a outra num lote com várias empresas - pior que simplesmente não achar.
    """
    normalized = normalize_cnpj(cnpj)
    if not normalized or not keys_folder.is_dir():
        return None
    if normalized in _only_digits(keys_folder.name):
        return keys_folder
    for candidate in keys_folder.iterdir():
        if candidate.is_dir() and normalized in _only_digits(candidate.name):
            return candidate
    return None


def find_keys_files_for_cnpj(keys_folder: Path, cnpj: str) -> list[Path]:
    """Localiza os arquivos de chaves NFC-e de uma empresa, aceitando três formatos:

    - Arquivo único: `keys_folder` é o próprio .xlsx, não uma pasta — ex. um arquivo
      solto baixado direto para Downloads, sem estrutura de pastas nenhuma. Devolvido
      direto, sem exigir nome/estrutura; quem garante que só entram chaves da empresa
      certa é o filtro por CNPJ embutido na chave, em `collect_nfce_keys_for_cnpj`.
    - Pasta de resultado do SIGA: dentro da subpasta '<COD> - <EMPRESA> - <CNPJ>',
      qualquer .xlsx é candidato, em qualquer profundidade/subpasta (mês, aba,
      sub-relatório...) — não é mais exigido um nome de pasta específico como
      'NFC-e', já que o SIGA nem sempre cria essa subpasta (ex.: a visualização
      "Emissor" de Informações Fiscais deixa o .xlsx solto direto na pasta do mês).
      Um .xlsx de NF-e/CT-e encontrado nessa varredura não "vaza" chave nenhuma: quem
      garante que só chaves de NFC-e entram é o filtro por modelo (`_is_nfce_key`,
      posições 21-22 = "65") em `extract_keys_from_spreadsheet`, aplicado por chave,
      não por pasta.
    - Formato antigo: um arquivo .xlsx solto direto em `keys_folder`, nomeado com o CNPJ.

    Devolve a união dos formatos aplicáveis, sem duplicar.
    """
    normalized = normalize_cnpj(cnpj)
    if not normalized:
        return []
    if keys_folder.is_file():
        return [keys_folder]
    if not keys_folder.is_dir():
        return []

    files: list[Path] = []
    company_folder = _company_folder_for_cnpj(keys_folder, cnpj)
    if company_folder is not None:
        files.extend(sorted(company_folder.rglob("*.xlsx")))

    for candidate in keys_folder.glob("*.xlsx"):
        if normalized in _only_digits(candidate.stem):
            files.append(candidate)

    seen_paths: set[Path] = set()
    unique_files: list[Path] = []
    for f in files:
        if f not in seen_paths:
            seen_paths.add(f)
            unique_files.append(f)
    return unique_files


def collect_nfce_keys_for_cnpj(keys_folder: Path, cnpj: str) -> dict[str, list[str]]:
    """Agrega, sem duplicar, as chaves NFC-e válidas de todos os arquivos encontrados
    para o CNPJ (arquivo único, pasta de resultado do SIGA, e/ou arquivo solto no
    formato antigo) — agrupadas por direção: `"saida"` (perfil "Emissor" no SIGA — a
    empresa emitiu), `"entrada"` (perfil "Destinatario" — a empresa recebeu) ou `""`
    quando o nome do arquivo não indica a direção (`tipo_nota_do_nome_arquivo`).

    A direção é uma propriedade do ARQUIVO de origem (cada arquivo do SIGA cobre só um
    perfil), não da chave em si — por isso é calculada uma vez por arquivo, não por chave.

    Sempre filtra pelo CNPJ do emitente embutido na própria chave (`cnpj_from_access_key`)
    antes de aceitar qualquer chave — não só quando `keys_folder` é um arquivo único (onde
    é a única forma de saber a quem a chave pertence), mas também no caso por pasta, como
    reforço extra: garante que uma chave de outra empresa nunca "vaza" para este resultado,
    mesmo que tenha sido encontrada por engano num arquivo/pasta com nome enganoso.
    """
    normalized = normalize_cnpj(cnpj)
    agrupado: dict[str, list[str]] = {"saida": [], "entrada": [], "": []}
    seen: set[str] = set()
    for path in find_keys_files_for_cnpj(keys_folder, cnpj):
        direcao = tipo_nota_do_nome_arquivo(path.name)
        for key in extract_keys_from_spreadsheet(path):
            if key in seen:
                continue
            if cnpj_from_access_key(key) != normalized:
                continue
            seen.add(key)
            agrupado[direcao].append(key)
    return agrupado


def extract_nfce_keys_from_text(texto: str) -> list[str]:
    """Extrai chaves de acesso de NFC-e (44 dígitos, modelo 65) de texto colado
    manualmente — uma por linha, ou separadas por espaço/vírgula/qualquer pontuação.
    Usado pela aba "Colar chaves" da interface, alternativa a apontar pasta/arquivo.
    """
    keys: list[str] = []
    seen: set[str] = set()
    for run in re.findall(r"\d+", texto or ""):
        if _is_nfce_key(run) and run not in seen:
            seen.add(run)
            keys.append(run)
    return keys


def group_nfce_keys_by_cnpj(keys: list[str]) -> dict[str, list[str]]:
    """Agrupa chaves já filtradas (modelo 65) pelo CNPJ do emitente embutido em cada
    uma — usado pelo modo de entrada manual (colar chaves), que não tem arquivo de
    origem para casar a chave com uma empresa por nome de pasta.
    """
    agrupado: dict[str, list[str]] = {}
    for chave in keys:
        agrupado.setdefault(cnpj_from_access_key(chave), []).append(chave)
    return agrupado


@dataclass(slots=True)
class NfceBaseInfo:
    """Uma linha da planilha-base CNPJ/IE, com um COD opcional (se a coluna existir)."""

    ie: str
    cod: str = ""


def load_cnpj_ie_base(path: Path | None) -> dict[str, NfceBaseInfo]:
    """Lê a planilha-base (colunas CNPJ/IE, e COD opcional) e devolve um mapa
    cnpj_normalizado -> NfceBaseInfo — a coluna COD é usada para nomear a pasta de saída
    no mesmo padrão "COD - EMPRESA - CNPJ" do modo SIGA (ver `run_batch_in_context`);
    sem essa coluna, `NfceBaseInfo.cod` fica vazio e o chamador cai em "SEM-COD".
    """
    if path is None or not path.exists():
        return {}
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        header = [str(cell.value or "").strip().upper() for cell in next(sheet.iter_rows(max_row=1))]
        try:
            cnpj_index = header.index("CNPJ")
            ie_index = header.index("IE")
        except ValueError:
            LOGGER.warning("Planilha-base NFC-e sem colunas CNPJ/IE: %s", path)
            workbook.close()
            return {}
        cod_index = header.index("COD") if "COD" in header else None

        mapa: dict[str, NfceBaseInfo] = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if len(row) <= max(cnpj_index, ie_index):
                continue
            cnpj = normalize_cnpj(row[cnpj_index])
            ie = normalize_ie(row[ie_index])
            cod = str(row[cod_index] or "").strip() if cod_index is not None and cod_index < len(row) else ""
            if cnpj:
                mapa[cnpj] = NfceBaseInfo(ie=ie, cod=cod)
        workbook.close()
        return mapa
    except Exception:  # noqa: BLE001
        LOGGER.exception("Falha ao ler a planilha-base de IE/CNPJ %s", path)
        return {}


@dataclass(slots=True)
class NfceCompanyLink:
    """Uma empresa disponível na área autenticada do portal, raspada após o login."""

    ie: str
    nome: str
    href: str  # script "javascript:..." que abre a aba de detalhe da empresa


@dataclass(slots=True)
class NfceBatchResult:
    """Resultado do processamento de uma linha (empresa) do lote NFC-e."""

    spreadsheet_row: SpreadsheetRow
    status: str  # "concluido" | "parcial" | "sem_chaves" | "sem_empresa" | "erro"
    chaves_total: int = 0
    chaves_baixadas: int = 0
    detalhe: str = ""


class NfceSessionManager:
    """Login, navegação e renovação de sessão no portal da SEFAZ-CE via selenium_compat."""

    def __init__(
        self,
        settings: Settings,
        page: Page,
        cancel_event: threading.Event | None = None,
        pause_event: threading.Event | None = None,
    ) -> None:
        self.settings = settings
        self.page = page
        self._login_time: float | None = None
        self._cancel_event = cancel_event
        self._pause_event = pause_event

    def _describe_unexpected_alert(self, exc: UnexpectedAlertError) -> str:
        """Fecha o alerta nativo que interrompeu o comando e monta uma mensagem clara.

        O alerta mais comum aqui e o "Tempo limite excedido! Sera necessario reiniciar a
        operacao" do proprio portal da SEFAZ-CE — ele aparece tanto por expiracao real de
        sessao quanto quando ja existe OUTRA sessao ativa com o mesmo CPF (o portal derruba
        a sessao nova). O codigo nao tem como diferenciar as duas causas a partir do texto
        do alerta sozinho, entao a mensagem cobre as duas.
        """
        alert_text = (getattr(exc, "alert_text", None) or "").strip()
        dismissed_text = (self.page.dismiss_alert_if_present() or "").strip()
        alert_text = alert_text or dismissed_text
        detalhe = f' ("{alert_text}")' if alert_text else ""
        return (
            f"O portal SEFAZ-CE encerrou a sessão inesperadamente{detalhe}. Isso costuma "
            "acontecer quando já existe outra sessão ativa com o mesmo CPF em outro "
            "computador ou navegador — feche a outra sessão e tente novamente."
        )

    def fazer_login(self) -> bool:
        """Efetua login com CPF/senha (settings.nfce_cpf/nfce_senha); idempotente se já logado.

        Quando o portal recusa o login por já existir outra sessão ativa com o mesmo CPF
        (ou um "Tempo limite excedido" equivalente — ver `_describe_unexpected_alert`),
        não desiste: espera `nfce_login_retry_wait_seconds` (padrão 30s) e tenta de novo,
        indefinidamente, até dar certo ou até o cancelamento cooperativo ser sinalizado.
        """
        # Checa a pagina ATUAL antes de navegar para qualquer lugar: se esta chamada
        # vem logo apos `discover_selectable_companies` (mesmo navegador, sessao ja
        # autenticada, parado na area de empresas), navegar de novo para
        # nfce_login_url tem comportamento imprevisivel no portal (o redirecionamento
        # pode nao cair nem na area de empresas nem na tela de login, derrubando as 3
        # tentativas abaixo com "elemento #txtUsuario nao encontrado"). Reaproveitar a
        # sessao pela pagina atual evita esse ponto de falha por completo.
        try:
            if self.page.locator("xpath=//a[contains(@href,'submete')]").count() > 0:
                LOGGER.info("Usuario NFC-e ja estava logado (sessao reaproveitada)")
                self._login_time = time.time()
                return True
        except Error:
            pass

        while True:
            resultado = self._tentar_login_3x()
            if resultado is True:
                return True
            if resultado is False:
                return False
            # `resultado` e a mensagem do alerta de sessao duplicada/expirada — espera e
            # tenta de novo, em vez de desistir (a causa costuma ser externa ao app: outra
            # sessao com o mesmo CPF que ainda nao foi fechada).
            if not self._aguardar_antes_de_retentar_login(resultado):
                return False

    def _tentar_login_3x(self) -> bool | str:
        """As 3 tentativas "normais" de login (falhas transitorias de rede/DOM).

        Devolve `True` (sucesso), `False` (as 3 tentativas esgotaram sem sucesso nem
        alerta) ou a mensagem do alerta de sessao duplicada/expirada (string) — nesse
        ultimo caso quem chama decide se espera e tenta de novo.
        """
        for tentativa in range(1, 4):
            try:
                self.page.goto(
                    self.settings.nfce_login_url,
                    wait_until="domcontentloaded",
                    timeout=self.settings.timeout_ms,
                )
                self.page.wait_for_timeout(1500)

                if self.page.locator("xpath=//a[contains(@href,'submete')]").count() > 0:
                    LOGGER.info("Usuario NFC-e ja estava logado")
                    self._login_time = time.time()
                    return True

                self.page.locator("#txtUsuario").fill(self.settings.nfce_cpf or "")
                self.page.locator("#txtSenha").fill(self.settings.nfce_senha or "")
                # Seleciona o perfil "CONTADOR" no dropdown de acesso (posição fixa no
                # formulário do portal, igual ao script original).
                self.page.evaluate(
                    """
                    () => {
                        const select = document.evaluate(
                            "/html/body/div/div[2]/div[2]/div/div[2]/form/table/tbody/tr[3]/td/select",
                            document, null, XPathResult.FIRST_ORDERED_NODE_TYPE, null
                        ).singleNodeValue;
                        if (!select) return;
                        const options = Array.from(select.options);
                        const match = options.find((o) => o.text.trim().toUpperCase() === "CONTADOR");
                        select.value = (match || options[1] || options[0]).value;
                        select.dispatchEvent(new Event("change"));
                    }
                    """
                )
                self.page.locator("#btEntrar").click()
                self.page.wait_for_timeout(4000)
            except UnexpectedAlertError as exc:
                message = self._describe_unexpected_alert(exc)
                LOGGER.warning("Alerta inesperado do portal durante o login NFC-e: %s", message)
                return message
            except Error as exc:
                LOGGER.warning("Tentativa %s de login NFC-e falhou: %s", tentativa, exc)
                self.page.wait_for_timeout(2000)
                continue

            if "erro" not in self.page.url.lower():
                LOGGER.info("Login NFC-e realizado com sucesso na tentativa %s", tentativa)
                self._login_time = time.time()
                return True

        return False

    def _aguardar_antes_de_retentar_login(self, motivo: str) -> bool:
        """Espera `nfce_login_retry_wait_seconds` de forma cancelável antes de retentar o login.

        Devolve `False` se o cancelamento foi sinalizado durante a espera (o chamador deve
        desistir); `True` para seguir e tentar o login de novo.
        """
        espera_seg = self.settings.nfce_login_retry_wait_seconds
        narrate_warning(
            "%s Tentando novamente em %s minuto(s)...",
            motivo,
            espera_seg // 60,
        )
        deadline = time.time() + espera_seg
        while time.time() < deadline:
            if not wait_if_paused(self._cancel_event, self._pause_event):
                narrate_warning("Retentativa de login cancelada pelo usuário.")
                return False
            time.sleep(1)
        return wait_if_paused(self._cancel_event, self._pause_event)

    def acessar_area_empresas(self) -> None:
        try:
            self.page.goto(
                self.settings.nfce_empresas_url,
                wait_until="domcontentloaded",
                timeout=self.settings.timeout_ms,
            )
            self.page.wait_for_timeout(1500)
        except UnexpectedAlertError as exc:
            message = self._describe_unexpected_alert(exc)
            LOGGER.warning("Alerta inesperado do portal ao acessar a área de empresas: %s", message)
            raise RuntimeError(message) from exc

    def autenticar(self) -> None:
        """Login + acesso à área de empresas, com a mesma retentativa indefinida do
        `fazer_login` também para o alerta de sessão duplicada/expirada aqui.

        Na prática, o alerta "Tempo limite excedido" costuma aparecer bem AQUI (ao
        navegar para a área de empresas logo após um login que já deu certo), não
        durante o preenchimento do formulário de login em si — então cobrir só
        `fazer_login` deixava esse caso escapando sem retry nenhum. Como a sessão
        pode já estar comprometida nesse ponto, refaz login + acesso do zero a cada
        tentativa, em vez de só repetir a navegação.
        """
        while True:
            if not self.fazer_login():
                raise RuntimeError("Falha no login automático do modo NFC-e.")
            try:
                self.acessar_area_empresas()
                return
            except RuntimeError as exc:
                if not self._aguardar_antes_de_retentar_login(str(exc)):
                    raise RuntimeError("Falha no login automático do modo NFC-e (cancelado pelo usuário).") from exc

    def sessao_expirando(self) -> bool:
        """True se o login já está velho o suficiente para renovar antes do portal derrubar."""
        if self._login_time is None:
            return True
        return (time.time() - self._login_time) >= self.settings.nfce_session_renewal_seconds

    def renovar_login(self) -> bool:
        """Reloga proativamente (logout + login) para evitar o "Tempo limite excedido" do portal."""
        narrate_warning("Renovando a sessão do portal SEFAZ-CE proativamente...")
        try:
            self.page.goto(
                self.settings.nfce_logout_url,
                wait_until="domcontentloaded",
                timeout=self.settings.timeout_ms,
            )
            self.page.wait_for_timeout(1500)
        except Error:
            pass

        if self.fazer_login():
            try:
                self.acessar_area_empresas()
            except RuntimeError as exc:
                # O portal pode recusar o acesso mesmo com o login OK (ex.: alerta de
                # sessão duplicada) — devolve False em vez de deixar a exceção propagar,
                # para que o chamador caia no protocolo de recuperação total em vez de
                # derrubar o lote inteiro.
                narrate_warning("Sessão renovada, mas o portal recusou o acesso à área de empresas: %s", exc)
                return False
            narrate_success("Sessão do portal SEFAZ-CE renovada com sucesso.")
            return True
        narrate_warning("Não foi possível renovar a sessão proativamente.")
        return False

    def reset_completo_com_espera(self) -> bool:
        """Protocolo de recuperação total: logout + espera longa + login do zero.

        A espera (~4 minutos por padrão) é intencional: dá tempo do portal liberar a
        sessão anterior por completo antes de tentar um novo login, evitando um segundo
        "Tempo limite excedido" imediato.
        """
        wait_minutes = self.settings.nfce_session_recovery_wait_seconds // 60
        narrate_warning(
            "Sessão do portal SEFAZ-CE instável; iniciando protocolo de recuperação "
            "(aguarde cerca de %s minuto(s), isso é esperado)...",
            wait_minutes,
        )
        try:
            self.page.goto(
                self.settings.nfce_logout_url,
                wait_until="domcontentloaded",
                timeout=self.settings.timeout_ms,
            )
        except Error:
            pass
        self.page.wait_for_timeout(self.settings.nfce_session_recovery_wait_seconds * 1000)

        if self.fazer_login():
            try:
                self.acessar_area_empresas()
            except RuntimeError as exc:
                narrate_error(
                    "Protocolo de recuperação: login OK, mas o portal recusou o acesso à área de empresas: %s", exc
                )
                return False
            narrate_success("Protocolo de recuperação concluído; sessão restabelecida.")
            return True
        narrate_error("Protocolo de recuperação falhou; sessão não pôde ser restabelecida.")
        return False

    def recuperar_sessao(self) -> bool:
        """Encadeia as 3 camadas de recuperação de sessão, da mais leve à mais persistente:
        relogin leve (`renovar_login`) → recuperação total com espera de 4min
        (`reset_completo_com_espera`) → se AINDA ASSIM falhar (sessão realmente travada
        por mais tempo do que essas duas dão conta), cai no mesmo retry indefinido de
        `autenticar()` (espera `nfce_login_retry_wait_seconds` e tenta de novo, sem
        limite, até dar certo ou até o cancelamento ser sinalizado).

        Sem essa 3ª camada, um lote longo (muitas empresas/chaves, várias horas) que
        esbarrasse numa instabilidade de sessão mais teimosa desistiria da empresa atual
        e de TODAS as seguintes — cada uma bateria na mesma sessão ainda quebrada logo
        na primeira checagem (`sessao_expirando`) e desistiria de novo, silenciosamente,
        pelo resto do lote.
        """
        if self.renovar_login() or self.reset_completo_com_espera():
            return True
        try:
            self.autenticar()
            return True
        except RuntimeError:
            return False

    def listar_empresas(self) -> list[NfceCompanyLink]:
        """Raspa a lista de empresas vinculadas ao CPF logado (IE + nome + link de abertura)."""
        try:
            raw = (
                self.page.evaluate(
                    """
                    () => Array.from(document.querySelectorAll("a[href*='submete']"))
                        .map((a) => ({ text: (a.innerText || "").trim(), href: a.getAttribute("href") || "" }))
                        .filter((item) => item.text)
                    """
                )
                or []
            )
        except UnexpectedAlertError as exc:
            message = self._describe_unexpected_alert(exc)
            LOGGER.warning("Alerta inesperado do portal ao listar as empresas: %s", message)
            raise RuntimeError(message) from exc

        agrupados: dict[str, dict[str, str]] = {}
        for item in raw:
            href = item["href"]
            texto = item["text"]
            registro = agrupados.setdefault(href, {"ie": "", "nome": ""})
            limpo = texto.replace(".", "").replace("-", "").replace("/", "")
            if limpo.isdigit():
                registro["ie"] = texto
            else:
                partes = texto.split(" ", 1)
                if len(partes) > 1 and partes[0].replace(".", "").replace("-", "").isdigit():
                    registro["ie"] = partes[0]
                    registro["nome"] = partes[1]
                else:
                    registro["nome"] = texto

        return [
            NfceCompanyLink(ie=dados["ie"] or "N/A", nome=dados["nome"] or "N/A", href=href)
            for href, dados in agrupados.items()
        ]


class NfceDownloadManager:
    """Detecta o XML baixado (via BrowserContext) e move para a pasta da empresa."""

    def __init__(self, context: BrowserContext, empresa_dir: Path) -> None:
        self.context = context
        self.empresa_dir = empresa_dir
        self.empresa_dir.mkdir(parents=True, exist_ok=True)

    def ja_baixado(self, chave: str) -> bool:
        """True se essa chave já foi baixada em execução anterior — busca recursiva na
        pasta inteira da empresa (não só na subpasta da direção atual), já que a mesma
        chave pode ter sido salva antes com uma classificação diferente de direção.
        Evita reconsultar o portal SEFAZ-CE por algo que já está no disco.
        """
        return any(self.empresa_dir.rglob(f"{chave}.*"))

    def mover_arquivo(self, origem: Path, chave: str, direcao: str = "") -> Path:
        # Garante que o nome final sempre carregue a chave, mesmo que o portal sirva
        # um nome genérico — facilita conferir depois qual XML corresponde a qual chave.
        nome_arquivo = origem.name if chave in origem.stem else f"{chave}{origem.suffix}"
        pasta_destino = self.empresa_dir / DIRECAO_SUBPASTA[direcao] if direcao else self.empresa_dir
        pasta_destino.mkdir(parents=True, exist_ok=True)
        destino = pasta_destino / nome_arquivo
        contador = 1
        while destino.exists():
            destino = pasta_destino / f"{destino.stem}_{contador}{destino.suffix}"
            contador += 1
        shutil.move(str(origem), str(destino))
        return destino


class NfceBatchExtractor:
    """Orquestra login, raspagem de empresas e download em lote de XMLs de NFC-e."""

    def __init__(
        self,
        settings: Settings,
        output_dir: Path,
        keys_folder: Path | None = None,
        base_spreadsheet_path: Path | None = None,
        manual_keys_text: str | None = None,
        manual_direcao: str = "",
    ) -> None:
        self.settings = settings
        self.output_dir = output_dir
        self.keys_folder = keys_folder
        self.base_spreadsheet_path = base_spreadsheet_path
        # Modo alternativo de entrada: chaves coladas manualmente em vez de pasta/arquivo
        # (aba "Colar chaves" da interface) — agrupadas por CNPJ uma única vez aqui, já
        # que não há arquivo de origem pra casar cada chave com uma empresa por nome de
        # pasta. `manual_direcao` ("saida"/"entrada"/"") classifica o lote inteiro de
        # uma vez, já que chave colada não carrega nome de arquivo pra detectar sozinha.
        self._manual_keys_by_cnpj = (
            group_nfce_keys_by_cnpj(extract_nfce_keys_from_text(manual_keys_text)) if manual_keys_text else None
        )
        self._manual_direcao = manual_direcao

    def _chaves_para_cnpj(self, cnpj: str) -> dict[str, list[str]]:
        """Resolve as chaves de uma empresa, priorizando o modo de entrada manual (se
        configurado) sobre a pasta/arquivo de chaves."""
        if self._manual_keys_by_cnpj is not None:
            chaves = self._manual_keys_by_cnpj.get(normalize_cnpj(cnpj), [])
            agrupado = {"saida": [], "entrada": [], "": []}
            agrupado[self._manual_direcao if self._manual_direcao in agrupado else ""] = chaves
            return agrupado
        if self.keys_folder is None:
            return {"saida": [], "entrada": [], "": []}
        return collect_nfce_keys_for_cnpj(self.keys_folder, cnpj)

    def run_batch_in_context(
        self,
        context: BrowserContext,
        selected_rows: list[SpreadsheetRow],
        cancel_event: threading.Event | None = None,
        pause_event: threading.Event | None = None,
    ) -> list[NfceBatchResult]:
        page = context.pages[0] if context.pages else context.new_page()
        session = NfceSessionManager(self.settings, page, cancel_event=cancel_event, pause_event=pause_event)

        narrate("Fazendo login no portal SEFAZ-CE (modo NFC-e)...")
        session.autenticar()
        narrate_success("Login no portal SEFAZ-CE realizado com sucesso.")

        narrate("Carregando lista de empresas vinculadas ao CPF...")
        empresas = session.listar_empresas()
        narrate_success("%s empresa(s) encontrada(s) na sessão.", len(empresas))

        mapa_cnpj_base = load_cnpj_ie_base(self.base_spreadsheet_path)
        mapa_ie_cnpj = {info.ie: cnpj for cnpj, info in mapa_cnpj_base.items() if info.ie}
        # Rede de segurança pro COD da pasta de saída quando a planilha-base CNPJ/IE
        # selecionada na tela não tiver essa coluna preenchida pra empresa (mesma
        # planilha-base COD/EMPRESA/CNPJ que o modo NF-e já usa para o mesmo fim).
        mapa_cod_empresa_cnpj = (
            load_cod_empresa_cnpj_base(self.settings.cod_empresa_cnpj_base_path)
            if self.settings.cod_empresa_cnpj_base_path
            else {}
        )

        results: list[NfceBatchResult] = []
        for index, spreadsheet_row in enumerate(selected_rows, start=1):
            if not wait_if_paused(cancel_event, pause_event):
                LOGGER.info("Lote NFC-e encerrado a pedido do usuario antes da proxima empresa.")
                narrate_warning("Execução encerrada pelo usuário.")
                break

            target_cnpj = normalize_cnpj(spreadsheet_row.cnpj)
            empresa_nome = spreadsheet_row.empresa or "SEM-EMPRESA"
            narrate(
                "Processando empresa %s de %s: %s (CNPJ %s)...",
                index,
                len(selected_rows),
                empresa_nome,
                spreadsheet_row.cnpj,
            )

            company = self._match_company(empresas, target_cnpj, mapa_ie_cnpj)
            if company is None:
                narrate_warning("Empresa %s não encontrada na sessão do portal SEFAZ-CE.", empresa_nome)
                results.append(NfceBatchResult(spreadsheet_row, status="sem_empresa"))
                continue

            chaves_por_direcao = self._chaves_para_cnpj(target_cnpj)
            chaves = [chave for lista in chaves_por_direcao.values() for chave in lista]
            if not chaves:
                narrate_warning(
                    "Nenhuma chave de NFC-e encontrada para %s na pasta configurada "
                    "(nem em subpasta 'NFC-e' de resultado do SIGA, nem em planilha solta).",
                    empresa_nome,
                )
                results.append(NfceBatchResult(spreadsheet_row, status="sem_chaves"))
                continue
            direcao_por_chave = {
                chave: direcao for direcao, lista in chaves_por_direcao.items() for chave in lista
            }

            narrate("%s chave(s) a processar para %s.", len(chaves), empresa_nome)
            # Mesmo padrão "COD - EMPRESA - CNPJ" do modo SIGA (ver _build_taxpayer_folder_name
            # em siga_extractor.py). Prioridade do COD: 1º a planilha-base CNPJ/IE selecionada
            # na tela; 2º a planilha-base COD/EMPRESA/CNPJ (mesma rede de segurança que o NF-e
            # já usa), pro caso de a empresa não ter COD preenchido na primeira; por fim
            # "SEM-COD" se nem isso resolver, igual ao fallback que o SIGA já usa.
            info_cod = mapa_cnpj_base.get(target_cnpj)
            cod = info_cod.cod if info_cod else ""
            if not cod:
                info_cod_extra = mapa_cod_empresa_cnpj.get(target_cnpj)
                cod = info_cod_extra.cod if info_cod_extra else ""
            cod = cod or "SEM-COD"
            empresa_dir = self.output_dir / sanitize_folder_name(f"{cod} - {empresa_nome} - {target_cnpj}")
            try:
                baixadas = self._processar_empresa(
                    context,
                    page,
                    session,
                    company,
                    chaves,
                    empresa_dir,
                    cancel_event,
                    pause_event,
                    direcao_por_chave,
                )
            except Exception as exc:  # noqa: BLE001
                # Uma falha imprevista (ex.: sessão instável que nem o protocolo de
                # recuperação total conseguiu restabelecer) não pode derrubar o lote
                # inteiro — registra esta empresa como erro e segue para a próxima.
                LOGGER.exception("Falha inesperada ao processar %s no modo NFC-e", empresa_nome)
                narrate_error("Falha inesperada ao processar %s: %s", empresa_nome, exc)
                results.append(NfceBatchResult(spreadsheet_row, status="erro", chaves_total=len(chaves)))
                continue

            status = "concluido" if baixadas == len(chaves) else "parcial"
            if status == "concluido":
                narrate_success("Empresa %s concluída: %s de %s chave(s) baixadas.", empresa_nome, baixadas, len(chaves))
            else:
                narrate_warning("Empresa %s incompleta: %s de %s chave(s) baixadas.", empresa_nome, baixadas, len(chaves))
            results.append(
                NfceBatchResult(
                    spreadsheet_row,
                    status=status,
                    chaves_total=len(chaves),
                    chaves_baixadas=baixadas,
                )
            )

        return results

    def discover_selectable_companies(
        self,
        context: BrowserContext,
        cancel_event: threading.Event | None = None,
        pause_event: threading.Event | None = None,
    ) -> list[dict]:
        """Loga no portal e devolve as empresas elegíveis para a tela de seleção.

        Elegível = a IE devolvida pelo portal resolve para um CNPJ na planilha-base
        IE/CNPJ *e* esse CNPJ tem um arquivo de chaves correspondente na pasta
        configurada (com pelo menos uma chave válida). Usado só para popular a grade
        antes da execução; `run_batch_in_context` refaz login/listagem do zero por
        conta própria (idempotente, já que `fazer_login` detecta sessão já ativa).
        """
        page = context.pages[0] if context.pages else context.new_page()
        session = NfceSessionManager(self.settings, page, cancel_event=cancel_event, pause_event=pause_event)

        narrate("Fazendo login no portal SEFAZ-CE (modo NFC-e)...")
        session.autenticar()
        narrate_success("Login no portal SEFAZ-CE realizado com sucesso.")

        empresas = session.listar_empresas()
        narrate("%s empresa(s) encontrada(s) na sessão do portal.", len(empresas))

        mapa_cnpj_base = load_cnpj_ie_base(self.base_spreadsheet_path)
        mapa_ie_cnpj = {info.ie: cnpj for cnpj, info in mapa_cnpj_base.items() if info.ie}

        rows: list[dict] = []
        for index, company in enumerate(empresas, start=1):
            cnpj = mapa_ie_cnpj.get(normalize_ie(company.ie), "")
            if not cnpj:
                continue
            total_chaves = sum(len(lista) for lista in self._chaves_para_cnpj(cnpj).values())
            if not total_chaves:
                continue
            rows.append(
                {
                    "row_number": index,
                    "cod": company.ie,
                    "empresa": company.nome,
                    "cnpj": cnpj,
                    "chaves_count": total_chaves,
                }
            )

        narrate_success("%s empresa(s) elegível(is) para download (com chaves disponíveis).", len(rows))
        return rows

    def _match_company(
        self,
        empresas: list[NfceCompanyLink],
        target_cnpj: str,
        mapa_ie_cnpj: dict[str, str],
    ) -> NfceCompanyLink | None:
        for company in empresas:
            ie_normalizada = normalize_ie(company.ie)
            if mapa_ie_cnpj.get(ie_normalizada) == target_cnpj:
                return company
        return None

    def _processar_empresa(
        self,
        context: BrowserContext,
        page: Page,
        session: NfceSessionManager,
        company: NfceCompanyLink,
        chaves: list[str],
        empresa_dir: Path,
        cancel_event: threading.Event | None = None,
        pause_event: threading.Event | None = None,
        direcao_por_chave: dict[str, str] | None = None,
    ) -> int:
        """Processa as chaves de uma empresa, com renovação de sessão e recuperação em falha.

        Preserva as três defesas do script original: renovação proativa por tempo
        (evita o "Tempo limite excedido" do portal), reset completo em erro crítico
        (sessão realmente caída) e limite de passadas para não entrar em loop infinito
        com chaves que nunca baixam (canceladas/denegadas/inexistentes na SEFAZ).
        """
        download_manager = NfceDownloadManager(context, empresa_dir)
        baixadas: set[str] = set()
        ja_baixadas = [chave for chave in chaves if download_manager.ja_baixado(chave)]
        if ja_baixadas:
            narrate(
                "%s chave(s) de %s já baixadas em execução anterior; pulando.",
                len(ja_baixadas),
                company.nome,
            )
            baixadas.update(ja_baixadas)
        pendentes = [chave for chave in chaves if chave not in baixadas]
        passadas_sem_progresso = 0
        direcao_por_chave = direcao_por_chave or {}

        for _passada in range(5):
            if not pendentes:
                break
            if not wait_if_paused(cancel_event, pause_event):
                LOGGER.info("Processamento de %s encerrado a pedido do usuario.", company.nome)
                break

            if session.sessao_expirando():
                if not session.recuperar_sessao():
                    break

            script = company.href.replace("javascript:", "")
            try:
                empresa_page = self._abrir_empresa(page, script)
            except (Error, TimeoutError) as exc:
                LOGGER.warning("Falha ao abrir a empresa %s: %s", company.nome, exc)
                if not session.recuperar_sessao():
                    break
                continue

            progresso_nesta_passada = 0
            falhas_consecutivas = 0
            try:
                for chave_index, chave in enumerate(list(pendentes), start=1):
                    if not wait_if_paused(cancel_event, pause_event):
                        LOGGER.info("Processamento de %s encerrado a pedido do usuario.", company.nome)
                        return len(baixadas)

                    if chave_index % self.settings.nfce_session_renewal_check_every_n_keys == 0 and session.sessao_expirando():
                        narrate("Renovando sessão antes de continuar o lote de %s...", company.nome)
                        break

                    if chave_index > 1:
                        # O portal SEFAZ-CE só exibe de forma confiável o botão de baixar XML
                        # na PRIMEIRA consulta feita depois de a página carregar — confirmado
                        # ao vivo (fora da automação, direto no navegador): a partir da 2ª
                        # consulta na mesma aba, o botão para de aparecer, mesmo com os dados
                        # da nota já buscados com sucesso pelo próprio portal (a chamada
                        # completa, o AngularJS recebe a resposta — só o popup não é exibido).
                        # Recarregar antes de cada chave (menos a 1ª, já fresca por causa de
                        # `_abrir_empresa`) evita esse travamento binário, em vez de só
                        # descobri-lo tarde via o gatilho de falhas consecutivas abaixo.
                        try:
                            empresa_page.reload(wait_until="domcontentloaded", timeout=self.settings.timeout_ms)
                            empresa_page.wait_for_timeout(1500)
                        except (Error, TimeoutError) as exc:
                            LOGGER.warning(
                                "Falha ao recarregar a página de consulta para %s: %s", company.nome, exc
                            )
                            # Segue tentando mesmo assim; se a consulta falhar por causa
                            # disso, cai no mesmo tratamento de falha consecutiva de sempre.

                    sucesso, precisa_reset = self._baixar_xml_chave(
                        context, empresa_page, chave, download_manager, direcao_por_chave.get(chave, "")
                    )
                    if sucesso:
                        baixadas.add(chave)
                        pendentes.remove(chave)
                        progresso_nesta_passada += 1
                        falhas_consecutivas = 0
                    else:
                        falhas_consecutivas += 1
                    if precisa_reset:
                        narrate_warning("Sessão instável durante o download de %s; acionando recuperação.", company.nome)
                        if not session.recuperar_sessao():
                            return len(baixadas)
                        break
                    if falhas_consecutivas >= 3:
                        # Mesmo gatilho do script original (erros_consecutivos >= 3 em
                        # processar_lote_chaves, importação/NFCE/codigo-fonte/xml nfce.py).
                        # Com o reload por chave acima, isso deixa de ser a defesa principal
                        # contra o botão não aparecer (causa raiz já coberta) e vira rede de
                        # segurança para outros problemas (sessão instável, erro de rede no
                        # próprio reload) — em vez de esperar a passada inteira terminar (até
                        # centenas de chaves), força a recuperação e reabre a empresa do
                        # zero assim que 3 falhas seguidas são detectadas.
                        narrate_warning(
                            "%s falha(s) seguida(s) processando %s; reabrindo a empresa.",
                            falhas_consecutivas,
                            company.nome,
                        )
                        if not session.recuperar_sessao():
                            return len(baixadas)
                        break
            finally:
                # Sempre fecha a aba da empresa ao fim da passada (sucesso, falha ou
                # reset) e volta pra lista — a próxima passada reabre a empresa do
                # zero, mesmo padrão do script original. Sem isso, cada passada
                # acumularia mais uma aba aberta.
                self._fechar_aba_empresa(empresa_page, page)

            if progresso_nesta_passada == 0:
                passadas_sem_progresso += 1
            else:
                passadas_sem_progresso = 0
            if passadas_sem_progresso >= 3:
                narrate_warning(
                    "%s chave(s) de %s não baixaram após tentativas repetidas "
                    "(provavelmente canceladas/denegadas/inexistentes na SEFAZ).",
                    len(pendentes),
                    company.nome,
                )
                break

        return len(baixadas)

    def _abrir_empresa(self, page: Page, script: str) -> Page:
        """Executa o script de abertura da empresa e navega para a consulta de NFC-e.

        O script (`javascript:submete(...)` raspado do portal) abre a empresa numa
        aba/janela NOVA via `window.open()` — nunca navega a aba atual. Precisa
        detectar essa aba nova (comparando os handles antes/depois) e devolvê-la,
        já que a aba original (lista de empresas) nunca vai ter o link "Consultar
        NFC-e". `launch_debug_browser` já abre o Chrome com `--disable-popup-blocking`
        para esse pop-up não ser bloqueado.
        """
        handles_antes = set(page.driver.window_handles)
        page.evaluate(script)
        page.wait_for_timeout(1500)

        novos_handles = [h for h in page.driver.window_handles if h not in handles_antes]
        if not novos_handles:
            raise TimeoutError(
                "A empresa não abriu uma aba nova (pop-up pode ter sido bloqueado pelo navegador)."
            )
        empresa_page = Page(page.context, novos_handles[0])

        empresa_page.evaluate(
            """
            () => {
                document.querySelectorAll('.mfe-migration-modal, .modal-backdrop, .modal').forEach((el) => el.remove());
                document.body.classList.remove('modal-open');
                document.body.style.overflow = 'auto';
            }
            """
        )
        link = empresa_page.locator("a[ui-sref='taxpayers.fiscalCouponsNfceList']")
        link.wait_for(state="visible", timeout=10_000)
        link.click(force=True)
        empresa_page.wait_for_timeout(2000)
        return empresa_page

    def _fechar_aba_empresa(self, empresa_page: Page, pagina_lista: Page) -> None:
        """Fecha a aba da empresa (aberta por `_abrir_empresa`) e volta para a lista."""
        driver = empresa_page.driver
        try:
            if empresa_page.handle in driver.window_handles:
                driver.switch_to.window(empresa_page.handle)
                driver.close()
        except Error:
            pass
        finally:
            if pagina_lista.handle in driver.window_handles:
                driver.switch_to.window(pagina_lista.handle)

    def _baixar_xml_chave(
        self,
        context: BrowserContext,
        page: Page,
        chave: str,
        download_manager: NfceDownloadManager,
        direcao: str = "",
    ) -> tuple[bool, bool]:
        """Consulta uma chave e baixa o XML. Retorna (sucesso, precisa_reset_de_sessao).

        Narra o resultado de cada chave (sucesso ou o motivo exato da falha) — antes,
        os quatro pontos de saída sem sucesso desta função eram todos silenciosos, e o
        operador só via "N chave(s) não baixaram" no resumo final, sem conseguir
        distinguir nota genuinamente inexistente/cancelada (nada a fazer) de nota
        encontrada mas cujo download travou (potencial instabilidade do portal).
        """
        try:
            page.evaluate(
                """
                (chave) => {
                    document.querySelectorAll('.modal-backdrop, .modal').forEach((el) => el.remove());
                    const input = document.getElementById('nfceKey');
                    if (input) {
                        input.value = chave;
                        input.dispatchEvent(new Event('input'));
                    }
                    const button = document.querySelector('button[ng-click="find()"]');
                    if (button) button.click();
                }
                """,
                chave,
            )

            try:
                resultados = page.locator("#table-search-coupons")
                resultados.wait_for(state="visible", timeout=10_000)
            except TimeoutError:
                narrate_warning("Chave %s: a consulta no portal não respondeu a tempo.", chave)
                return False, False

            achou_link = page.evaluate(
                """
                () => {
                    const rows = document.querySelectorAll('#table-search-coupons tbody tr');
                    if (rows.length === 0 || (rows.length === 1 && rows[0].innerText.includes('Nenhum registro'))) return false;
                    const link = rows[0].querySelector('a');
                    if (link) { link.click(); return true; }
                    return false;
                }
                """
            )
            if not achou_link:
                narrate_warning(
                    "Chave %s: não encontrada na consulta (nota cancelada/denegada/inexistente na SEFAZ, ou ainda não processada).",
                    chave,
                )
                return False, False

            try:
                botao_xml = page.locator("button[ng-click='downloadXML()']")
                # 20s (era 10s): consultas sucessivas na mesma aba de empresa (sem recarregar
                # a página) tendem a renderizar o botão mais devagar a cada nova chave — 10s
                # só bastava de forma confiável para a primeira consulta após abrir a empresa.
                botao_xml.wait_for(state="visible", timeout=20_000)
            except TimeoutError:
                narrate_warning("Chave %s: nota encontrada, mas o botão de baixar XML não apareceu a tempo.", chave)
                return False, False

            snapshot = context.download_snapshot()
            botao_xml.click(force=True)
            try:
                downloaded_path = context.wait_for_download(snapshot, timeout_ms=15_000)
            except TimeoutError:
                narrate_warning("Chave %s: nota encontrada e download clicado, mas o arquivo não chegou a tempo.", chave)
                return False, False

            download_manager.mover_arquivo(downloaded_path, chave, direcao)

            page.evaluate(
                """
                () => {
                    document.querySelectorAll('.modal-backdrop, .modal').forEach((el) => el.remove());
                }
                """
            )
            narrate_success("Chave %s baixada.", chave)
            return True, False
        except TimeoutError:
            narrate_warning("Chave %s: tempo esgotado ao consultar no portal.", chave)
            return False, False
        except Error as exc:
            texto_erro = str(exc).lower()
            precisa_reset = "disconnected" in texto_erro or "stacktrace" in texto_erro or "script timeout" in texto_erro
            LOGGER.warning("Falha ao baixar a chave %s: %s", chave, exc)
            narrate_warning("Chave %s: erro técnico ao processar (%s).", chave, exc)
            return False, precisa_reset
